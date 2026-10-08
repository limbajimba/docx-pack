"""Smoke tests for the pack scripts. Run: python3 tests/test_pack.py  (or python3 -m pytest tests/ -q)"""
import json, os, subprocess, sys, tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
EXAMPLE = ROOT / "examples" / "acme-memo.docx"
sys.path.insert(0, str(SCRIPTS))


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


def make_docx(paragraphs, bullets=(), heading=None):
    from docx import Document
    d = Document()
    if heading:
        d.add_heading(heading, 1)
    for t in paragraphs:
        d.add_paragraph(t)
    for lead, rest in bullets:
        p = d.add_paragraph(style="List Bullet")
        r = p.add_run(lead); r.bold = True
        p.add_run(rest)
    f = tempfile.NamedTemporaryFile(delete=False, suffix=".docx").name
    d.save(f)
    return f


def test_codify_reads_page_and_font():
    out = tempfile.mkdtemp()
    r = run(SCRIPTS / "codify_template.py", EXAMPLE, "--out", out)
    assert r.returncode == 0, r.stderr
    spec = json.load(open(Path(out) / "style-spec.json"))
    assert spec["page"]["size"] == "A4"
    assert spec["roles_inferred"]["font_body"] == "Gill Sans MT"
    assert spec["header_footer"]["headers"] and spec["header_footer"]["headers"][0]["images"]
    assert (Path(out) / "STYLE.md").exists()


def test_lint_flags_tells():
    f = make_docx(
        ["The purpose of this document is to delve into the landscape — it's not a bug, it's a feature. 🚀",
         "Furthermore, this is a testament to robust synergy."],
        bullets=[("Speed: ", "fast"), ("Cost: ", "cheap"), ("Risk: ", "low")],
        heading="Key Findings And Next Steps",
    )
    r = run(SCRIPTS / "lint_docx.py", f, "--json")
    assert r.returncode == 1  # emoji is a FAIL
    codes = {x["code"] for x in json.loads(r.stdout)["findings"]}
    for c in ("emoji", "em-dash", "ai-vocabulary", "not-x-but-y", "bold-lead-bullets", "purpose-opener", "title-case-headings"):
        assert c in codes, c


def test_lint_passes_clean_example():
    r = run(SCRIPTS / "lint_docx.py", EXAMPLE, "--spec", ROOT / "house" / "silvertree.style-spec.json", "--json")
    assert r.returncode == 0, r.stdout
    sev = {x["severity"] for x in json.loads(r.stdout)["findings"]}
    assert "FAIL" not in sev and "WARN" not in sev


def test_set_metadata_roundtrip():
    f = make_docx(["Hello."])
    r = run(SCRIPTS / "set_metadata.py", f, "--creator", "Acme Co", "--company", "Acme Co", "--title", "T")
    assert r.returncode == 0, r.stderr
    from docx import Document
    d = Document(f)
    assert d.core_properties.author == "Acme Co" and d.core_properties.last_modified_by == "Acme Co" and d.core_properties.title == "T"


def test_build_dedupes_callout_labels_and_sets_metadata():
    import shutil
    d = Path(tempfile.mkdtemp())
    content = json.load(open(ROOT / "examples" / "acme-content.json"))
    content["blocks"] = [b for b in content["blocks"] if b["type"] != "image"]
    for b in content["blocks"]:
        if b["type"] == "finding":
            b["counter"] = "Counter: " + b["counter"]
    json.dump(content, open(d / "c.json", "w"))
    shutil.copy(ROOT / "house" / "silvertree-logo.png", d / "silvertree-logo.png")
    r = subprocess.run(["node", str(SCRIPTS / "build_docx.js"), "--spec", str(ROOT / "house" / "silvertree.style-spec.json"), "--content", str(d / "c.json"), "--out", str(d / "o.docx")], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    from docx import Document
    doc = Document(str(d / "o.docx"))
    texts = [p.text for p in doc.paragraphs if p.text.startswith("Counter:")]
    assert texts and all(not t.startswith("Counter: Counter:") for t in texts)
    assert doc.core_properties.author == "SilverTree Equity"
    import zipfile
    assert b"<Application>" in zipfile.ZipFile(str(d / "o.docx")).read("docProps/app.xml")


def test_fill_template_uses_template_styles():
    from docx import Document
    t = make_docx(["template body"])
    out = tempfile.NamedTemporaryFile(delete=False, suffix=".docx").name
    r = run(SCRIPTS / "fill_template.py", "--template", t, "--content", ROOT / "examples" / "acme-content.json", "--out", out)
    assert r.returncode == 0, r.stderr
    d = Document(out)
    texts = [p.text for p in d.paragraphs]
    assert "template body" not in texts
    assert any(p.style.name == "Heading 1" for p in d.paragraphs)
    assert len(d.tables) >= 3


def test_landscape_figure_grid():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    d = Path(tempfile.mkdtemp())
    for n in (1, 2):
        fig, ax = plt.subplots(figsize=(4, 2)); ax.plot([1, 2], [n, 1]); fig.savefig(d / f"c{n}.png"); plt.close(fig)
    content = {"meta": {"title": "T", "kicker": "K", "date": "1 Jan 2026"}, "blocks": [
        {"type": "h1", "text": "Page one"},
        {"type": "figure_grid", "cols": 2, "cells": [
            {"path": str(d / "c1.png"), "title": "One", "sources": "Sources: F1", "note": "n"},
            {"path": str(d / "c2.png"), "title": "Two", "sources": "Sources: F2"}]}]}
    (d / "c.json").write_text(json.dumps(content))
    out = d / "l.docx"
    r = subprocess.run(["node", str(SCRIPTS / "build_docx.js"), "--spec", str(ROOT / "house" / "silvertree-landscape.style-spec.json"),
                        "--content", str(d / "c.json"), "--out", str(out)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    import zipfile
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert 'w:orient="landscape"' in xml and 'w:w="16838"' in xml and 'w:h="11906"' in xml, xml[-600:]
    from docx import Document
    doc = Document(str(out))
    cells = [c for t in doc.tables for row in t.rows for c in row.cells]
    assert sum(1 for c in cells if c._tc.xpath(".//w:drawing")) == 2



def test_chart_bars_tick_step():
    import chart_bars
    assert chart_bars.x_ticks({"max": 10}) == [0, 2, 4, 6, 8, 10]
    assert chart_bars.x_ticks({"max": 600, "step": 100}) == [0, 100, 200, 300, 400, 500, 600]


def test_render_preview_accepts_pptx():
    import shutil
    if not (shutil.which("soffice") and shutil.which("pdftoppm")):
        print("skip: soffice or pdftoppm missing"); return
    from pptx import Presentation
    d = Path(tempfile.mkdtemp())
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = "One slide"
    prs.save(d / "deck.pptx")
    r = subprocess.run(["bash", str(SCRIPTS / "render_preview.sh"), str(d / "deck.pptx"), str(d / "prev"), "40"],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    assert (d / "prev" / "deck.pdf").exists()
    assert "pages: 1" in r.stdout

if __name__ == "__main__":  # plain runner when pytest is not installed
    import traceback
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn(); print(f"ok   {name}")
            except Exception:
                fails += 1; print(f"FAIL {name}"); traceback.print_exc()
    sys.exit(1 if fails else 0)
